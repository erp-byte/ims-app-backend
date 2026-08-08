"""
Swap RECON22JUL placeholder boxes for the REAL scanned boxes that later arrived for the same
stock. DRY-RUN BY DEFAULT.

This is the recurring failure of the 22-Jul reconciliation: it minted synthetic RC22JUL-* rows
for stock that was physically present but not yet digitized. When the warehouse later scans
that stock for real, the real boxes land with lot_no NULL and the placeholders keep standing,
so the lot silently doubles. Already seen on 14570, 13766, 8130 and now 13865-13872.

Given a transaction holding the real NULL-lot boxes and the lots they belong to, this:
    1. asserts sum(sheet cartons for those lots) == count(NULL-lot boxes on that transaction)
       -- if the two do not agree exactly, nothing is written;
    2. allocates the real boxes to the lots in ascending numeric box_id suffix, lots in
       ascending order (nothing in the data distinguishes one identical box from another, so a
       deterministic rule is the only defensible choice);
    3. stamps the sheet's descriptive fields on the relabelled rows, preserving box_id and
       transaction_no so the sticker and inward link survive;
    4. deletes exactly the RECON22JUL placeholders for those lots, writing a
       cold_stock_disposition row with the full prior row in snapshot_data.

Rows are selected by transaction_no + NUMERIC suffix, never `box_id LIKE 'prefix%'` -- the
8-digit prefix is last-8-of-epoch-ms and recycles about every 27.8 hours.

Run:
  python scripts/swap_synthetic_for_real.py --txn TR-20260716141141 --company CFPL \
      --lots 13865,13866,13867,13868,13869,13870,13871,13872 --dry-run
"""
from __future__ import annotations
import os
import sys
import json
import argparse
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import psycopg
from dotenv import load_dotenv

sys.path.insert(0, str(Path(__file__).resolve().parent))
from report_lots_in_closing_29jul import norm, _s, COL, HEADER_ROW, PREFIX
from apply_reconcile_29jul import SHEET_COL_IDX, SHEET_TO_COLD, coerce, REF_SOURCES

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env")
EXCEL_PATH = ROOT / "data" / "Cold Storage Stock Details 29th July Closing.xlsx"
ACTOR = os.environ.get("RECON_ACTOR", "ai.1@candorfoods.in")


def load_sheet(lots):
    ws = openpyxl.load_workbook(str(EXCEL_PATH), read_only=True, data_only=True)["Sheet1"]
    want = {norm(l) for l in lots}
    out = {}
    for r in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        k = norm(_s(r[COL["lot_no"]]))
        if k not in want:
            continue
        d = out.setdefault(k, {"cart": 0.0, "fields": {},
                               "company": _s(r[COL["company"]]).upper()})
        d["cart"] += float(r[COL["cartons"]] or 0)
        for skey, (dbcol, kind) in SHEET_TO_COLD.items():
            v = coerce(kind, r[SHEET_COL_IDX[skey]])
            if v is not None:
                d["fields"].setdefault(dbcol, v)
    for d in out.values():
        d["cart"] = round(d["cart"])
    return out


def in_flight(cur, box_ids):
    hits = {}
    for tbl, cols in REF_SOURCES.items():
        for c in cols:
            cur.execute(f"SELECT count(*) FROM {tbl} WHERE {c}::text = ANY(%s)", (box_ids,))
            n = cur.fetchone()[0]
            if n:
                hits[f"{tbl}.{c}"] = n
    return hits


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--txn", required=True)
    ap.add_argument("--company", required=True, choices=list(PREFIX))
    ap.add_argument("--lots", required=True, help="comma-separated, ascending allocation order")
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--dry-run", action="store_true")
    g.add_argument("--execute", action="store_true")
    args = ap.parse_args()

    lots = [l.strip() for l in args.lots.split(",") if l.strip()]
    lots.sort(key=lambda l: (len(l), l))
    tbl = f"{PREFIX[args.company]}_cold_stocks"
    batch = f"SWAP-{args.txn}"
    sheet = load_sheet(lots)

    with psycopg.connect(os.environ["DATABASE_URL"]) as conn:
        cur = conn.cursor()
        cur.execute(f"""SELECT id, box_id FROM {tbl}
                        WHERE transaction_no = %s AND lot_no IS NULL
                        ORDER BY NULLIF(split_part(box_id, '-', 2), '')::int""", (args.txn,))
        real = cur.fetchall()
        want = sum(sheet.get(norm(l), {}).get("cart", 0) for l in lots)

        print(f"=== {batch} ({'EXECUTE' if args.execute else 'DRY RUN'}) ===\n")
        print(f"  real NULL-lot boxes on {args.txn}: {len(real)}")
        print(f"  sheet cartons across {len(lots)} lots      : {want}")
        if len(real) != want:
            raise SystemExit(f"ABORT: {len(real)} real boxes but the sheet wants {want}. "
                             f"Refusing to guess an allocation.")

        # allocate ascending suffix -> lots in ascending order
        plan, i = [], 0
        for l in lots:
            k = norm(l)
            n = sheet.get(k, {}).get("cart", 0)
            chunk = real[i:i + n]
            i += n
            cur.execute(f"""SELECT id, box_id FROM {tbl}
                            WHERE lot_no::text = %s AND transaction_no = 'RECON22JUL'""", (l,))
            syn = cur.fetchall()
            plan.append({"lot": l, "n": n, "real": chunk, "syn": syn,
                         "fields": sheet.get(k, {}).get("fields", {}),
                         "company": sheet.get(k, {}).get("company")})
            print(f"  lot {l:>7s}: {n:>4d} real {chunk[0][1]} .. {chunk[-1][1]}"
                  f"   replaces {len(syn):>4d} synthetic"
                  + ("" if len(syn) == n else f"   !! synthetic count != {n}")
                  + ("" if sheet.get(k, {}).get("company") == args.company
                     else f"   !! sheet company is {sheet.get(k, {}).get('company')}"))

        bad = [p for p in plan if len(p["syn"]) != p["n"]
               or p["company"] != args.company]
        hits = in_flight(cur, [b for _i, b in real]
                         + [b for p in plan for _i, b in p["syn"]])
        if hits:
            print(f"\n  !! in-flight references: {hits}")
        if bad:
            print(f"\n  !! {len(bad)} lots fail the synthetic-count / company check")
        print(f"\n  net rows in {tbl}: "
              f"{len(real) - sum(len(p['syn']) for p in plan):+d} "
              f"({len(real)} relabelled in place, "
              f"{sum(len(p['syn']) for p in plan)} placeholders deleted)")

        if not args.execute:
            conn.rollback()
            print("\nDRY RUN - nothing written.")
            return
        if hits or bad:
            raise SystemExit("ABORT: unsafe preconditions; nothing written.")

        backup = {"batch": batch, "deleted": [], "relabelled": []}
        for p in plan:
            syn_ids = [i for i, _b in p["syn"]]
            if syn_ids:
                cur.execute(f"SELECT to_jsonb(t) FROM {tbl} t WHERE id = ANY(%s)", (syn_ids,))
                backup["deleted"].extend(r[0] for r in cur.fetchall())
                cur.execute(f"""
                    INSERT INTO cold_stock_disposition
                      (box_id, transaction_no, lot_no, item_description, from_company, unit,
                       from_site, source_table, disposition_type, disposition_ref_table,
                       disposition_ref_no, disposed_by, snapshot_data, notes)
                    SELECT t.box_id, t.transaction_no, %s, t.item_description, %s, t.unit,
                           t.storage_location, %s, 'manual_correction', %s, %s, %s,
                           to_jsonb(t), %s
                    FROM {tbl} t WHERE t.id = ANY(%s)""",
                            (p["lot"], PREFIX[args.company], tbl, tbl, batch, ACTOR,
                             f"{batch}: RECON22JUL placeholder replaced by the real scanned "
                             f"boxes of {args.txn}", syn_ids))
                cur.execute(f"DELETE FROM {tbl} WHERE id = ANY(%s)", (syn_ids,))

            f = dict(p["fields"])
            wt, rate = f.get("weight_kg"), f.get("last_purchase_rate")
            f["total_inventory_kgs"] = wt
            f["value"] = wt * rate if wt is not None and rate is not None else None
            f["no_of_cartons"] = Decimal("1")
            ids = [i for i, _b in p["real"]]
            cur.execute(f"UPDATE {tbl} SET lot_no = %s, "
                        + ", ".join(f"{c} = %s" for c in f)
                        + ", updated_at = CURRENT_TIMESTAMP WHERE id = ANY(%s)",
                        [p["lot"]] + list(f.values()) + [ids])
            backup["relabelled"].append({"lot": p["lot"],
                                         "box_ids": [b for _i, b in p["real"]]})

        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = ROOT / f"swap_{args.txn}_backup_{stamp}.json"
        path.write_text(json.dumps(backup, indent=1, default=str))
        conn.commit()
        print(f"\nCOMMITTED. Backup: {path}")
        print(f"Reversible via cold_stock_disposition WHERE disposition_ref_no = '{batch}'.")


if __name__ == "__main__":
    main()
